from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = 'Microsoft YaHei'
font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

for i in range(1, 5):
    hs = doc.styles[f'Heading {i}']
    hs.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    hs.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

def add_code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x2D, 0x2D, 0x2D)
    from docx.oxml import OxmlElement
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F5F5F5')
    run.element.rPr.append(shd)

def add_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

def add_image_placeholder(doc, desc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'[ 插图：{desc} ]')
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

# ========== 封面 ==========
doc.add_paragraph()
doc.add_paragraph()
title = doc.add_heading('OpenCode 技术使用文档', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph('Olares 平台部署版')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(14)
subtitle.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
date_p = doc.add_paragraph('2026 年 3 月')
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
date_p.runs[0].font.size = Pt(12)
date_p.runs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)

doc.add_page_break()

# ========== 目录占位 ==========
doc.add_heading('目录', level=1)
doc.add_paragraph('（粘贴至飞书后可手动生成目录）')
doc.add_page_break()

# ========== 1. 概述 ==========
doc.add_heading('1. 概述', level=1)
doc.add_paragraph(
    'OpenCode 是一款 AI 编码助手，支持通过自然语言对话驱动代码编写、环境搭建、'
    '项目构建与网页预览等全流程开发工作。本文档基于 Olares 平台的容器化部署版本，'
    '涵盖安装架构、模型配置、包管理、Web 预览、插件系统及已知问题等内容。'
)
doc.add_paragraph(
    '官方文档：https://opencode.ai/docs'
)

# ========== 2. 安装架构 ==========
doc.add_heading('2. 安装架构', level=1)

doc.add_heading('2.1 容器架构', level=2)
doc.add_paragraph(
    'OpenCode 在 Olares 平台上以 Kubernetes Deployment 形式部署，包含以下容器：'
)
table = doc.add_table(rows=5, cols=3, style='Light List Accent 1')
headers = ['容器', '运行用户', '职责']
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
data = [
    ['init-setup', 'root (uid 0)', '初始化目录结构、写入 Skill 文件与配置'],
    ['init-packages', 'root (uid 0)', '安装预装软件包、创建文件系统快照'],
    ['opencode', 'opencode (uid 1000)', '运行 OpenCode Web UI 主服务'],
    ['pkg-manager', 'root (uid 0)', '后台包管理 sidecar，处理 pkg-install 请求'],
]
for row_idx, row_data in enumerate(data, 1):
    for col_idx, val in enumerate(row_data):
        table.rows[row_idx].cells[col_idx].text = val

doc.add_paragraph()
add_note(doc, '注：主容器以普通用户（uid 1000）运行，保障安全性；包管理操作通过 sidecar 以 root 权限代理执行。')

doc.add_heading('2.2 预装软件', level=2)
doc.add_paragraph('以下工具在应用安装时自动部署，无需额外安装：')
table2 = doc.add_table(rows=7, cols=2, style='Light List Accent 1')
table2.rows[0].cells[0].text = '类别'
table2.rows[0].cells[1].text = '工具'
preinstalled = [
    ['编程语言', 'Python 3, pip, Node.js, npm, Go, Rust, Cargo'],
    ['版本控制', 'Git'],
    ['网络工具', 'curl, wget, SSH, ss'],
    ['Shell', 'Bash'],
    ['压缩工具', 'zip, unzip, gzip, bzip2, xz, tar, zstd'],
    ['网络诊断', 'iproute2'],
]
for row_idx, (cat, tools) in enumerate(preinstalled, 1):
    table2.rows[row_idx].cells[0].text = cat
    table2.rows[row_idx].cells[1].text = tools

doc.add_heading('2.3 数据目录', level=2)
doc.add_paragraph('应用数据和用户工程文件分别挂载在不同路径：')
table3 = doc.add_table(rows=3, cols=3, style='Light List Accent 1')
table3.rows[0].cells[0].text = '用途'
table3.rows[0].cells[1].text = '容器内路径'
table3.rows[0].cells[2].text = '文件系统位置'
table3.rows[1].cells[0].text = '配置与应用数据'
table3.rows[1].cells[1].text = '/home/opencode'
table3.rows[1].cells[2].text = 'Data/opencode'
table3.rows[2].cells[0].text = '用户工程文件'
table3.rows[2].cells[1].text = '/home/opencode/workspace'
table3.rows[2].cells[2].text = 'Home/Code'

# ========== 3. 模型配置 ==========
doc.add_heading('3. 模型配置', level=1)

doc.add_heading('3.1 默认模型', level=2)
doc.add_paragraph(
    'OpenCode 默认提供若干免费模型。启动后在界面左侧选择模型即可开始对话和编码。'
)
add_image_placeholder(doc, '默认模型选择界面截图')

doc.add_heading('3.2 连接第三方供应商', level=2)
doc.add_paragraph(
    '点击界面左下角的齿轮图标或"连接提供商"按钮，进入提供商设置页面。'
    '选择目标提供商和连接方式，输入 API URL 或 Key 即可连接。'
    '连接成功后，模型选择界面会出现该提供商的可用模型。'
)
add_image_placeholder(doc, '提供商配置界面截图')

doc.add_heading('3.3 连接 Olares Ollama', level=2)
doc.add_paragraph(
    '连接 Olares 自带的 Ollama 服务时，需使用「自定义供应商」方式，'
    '而非默认提供的 Ollama Cloud 选项。配置要点：'
)
doc.add_paragraph('基础 URL 必须以 /v1 结尾', style='List Bullet')
doc.add_paragraph('只能使用 Ollama API 入口（设置中的 Ollama API 地址），不能使用共享入口', style='List Bullet')
doc.add_paragraph(
    '在 AMD64 机器上安装了应用 1.0.1 版及以上的，也可以使用共享入口。',
    style='List Bullet'
)
add_image_placeholder(doc, 'Ollama 自定义供应商配置截图')

doc.add_heading('3.4 配置文件', level=2)
doc.add_paragraph(
    '自定义供应商的配置文件位于：'
)
add_code_block(doc, 'Data/opencode/.config/opencode/opencode.jsonc')
doc.add_paragraph(
    '该文件默认扩展名为 .jsonc，可改为 .json 后在 Olares 文件系统中直接编辑。'
    '编辑后无需改回 .jsonc，重启 OpenCode 应用即可生效。'
)
doc.add_paragraph(
    '配置文件格式详见官方文档：https://opencode.ai/docs/providers/#custom-provider'
)

# ========== 4. 包管理 ==========
doc.add_heading('4. 包管理（pkg-install）', level=1)

doc.add_heading('4.1 简介', level=2)
doc.add_paragraph(
    'pkg-install 是对 Alpine Linux apk 包管理器的封装命令。'
    '由于主容器以普通用户运行，无法直接使用 sudo 或 apk，'
    '所有系统级包管理操作统一通过 pkg-install 完成。'
)

doc.add_heading('4.2 命令参考', level=2)
table4 = doc.add_table(rows=8, cols=2, style='Light List Accent 1')
table4.rows[0].cells[0].text = '命令'
table4.rows[0].cells[1].text = '说明'
cmds = [
    ['pkg-install <包名> [包名2 ...]', '安装一个或多个包'],
    ['pkg-install <包名>=<版本>', '安装指定版本'],
    ['pkg-install --remove <包名>', '卸载包'],
    ['pkg-install --search <关键词>', '搜索可用包'],
    ['pkg-install --info <包名>', '查看包详情和可用版本'],
    ['pkg-install --list', '列出所有已安装的包'],
    ['pkg-install --help', '显示帮助信息'],
]
for i, (cmd, desc) in enumerate(cmds, 1):
    table4.rows[i].cells[0].text = cmd
    table4.rows[i].cells[1].text = desc

doc.add_heading('4.3 常用包速查', level=2)
table5 = doc.add_table(rows=11, cols=2, style='Light List Accent 1')
table5.rows[0].cells[0].text = '需求'
table5.rows[0].cells[1].text = '包名'
pkgs = [
    ['Java 17 / 21', 'openjdk17-jdk / openjdk21-jdk'],
    ['GCC / G++ / Make', 'build-base'],
    ['CMake', 'cmake'],
    ['FFmpeg', 'ffmpeg'],
    ['PostgreSQL 客户端', 'postgresql-client'],
    ['MySQL 客户端', 'mariadb-client'],
    ['PHP', 'php83'],
    ['Ruby / Perl / Lua', 'ruby / perl / lua5.4'],
    ['.NET SDK', 'dotnet8-sdk'],
    ['vim / neovim / tmux', 'vim / neovim / tmux'],
]
for i, (need, pkg) in enumerate(pkgs, 1):
    table5.rows[i].cells[0].text = need
    table5.rows[i].cells[1].text = pkg

doc.add_heading('4.4 语言包管理器', level=2)
doc.add_paragraph('对于语言生态内的包，直接使用原生工具，无需 pkg-install：')
add_code_block(doc, 'pip install <pkg>           # Python\nnpm install <pkg>           # Node.js\ngo install <pkg>@latest     # Go\ncargo install <pkg>         # Rust')

doc.add_heading('4.5 通过 AI 对话管理包', level=2)
doc.add_paragraph(
    '除了在终端中直接使用 pkg-install 外，还可以通过 AI 对话指令让 OpenCode 安装或卸载包。'
    '例如：「请帮我安装 ffmpeg」。AI 会自动触发 system-admin Skill 并调用 pkg-install 完成操作。'
)
add_image_placeholder(doc, 'AI 对话安装包示意截图')
doc.add_paragraph(
    '若 AI 未自动触发 Skill，可手动加载：'
)
add_code_block(doc, '/skill load system-admin')

doc.add_heading('4.6 包持久化说明', level=2)
doc.add_paragraph(
    '通过 pkg-install 安装的包在同版本内持久化（重启不丢失）。'
    '版本升级时，系统会自动记录用户安装过的包，并在新版本快照完成后自动重新安装。'
)

# ========== 5. Web 预览 ==========
doc.add_heading('5. Web 预览', level=1)

doc.add_heading('5.1 工作原理', level=2)
doc.add_paragraph(
    '环境内置反向代理，容器中运行的开发服务器可通过 /__preview/<端口>/ 路径在浏览器中访问。'
    '无需额外配置域名或端口映射。'
)

doc.add_heading('5.2 使用方法', level=2)
doc.add_paragraph('通过 AI 对话指令启动预览。例如：「请启动本文件夹的工程网页，端口用 4000」。')
doc.add_paragraph('AI 会自动触发 web-preview Skill，执行以下操作：')
doc.add_paragraph('启动开发服务器（绑定 0.0.0.0，设置 base path）', style='List Number')
doc.add_paragraph('验证服务器运行状态', style='List Number')
doc.add_paragraph('生成预览 URL 并返回给用户', style='List Number')

doc.add_paragraph()
doc.add_paragraph('预览地址格式：')
add_code_block(doc, 'https://<你的 OpenCode 域名>/__preview/<端口>/')
doc.add_paragraph(
    '其中域名即为浏览器地址栏中访问 OpenCode 的域名。'
)
add_image_placeholder(doc, 'Web 预览效果截图')
doc.add_paragraph('若 AI 未自动触发 Skill，可手动加载：')
add_code_block(doc, '/skill load web-preview')

# ========== 6. Skill 系统 ==========
doc.add_heading('6. Skill 系统', level=1)

doc.add_heading('6.1 概述', level=2)
doc.add_paragraph(
    'Skill 是 OpenCode 的能力扩展模块，以 Markdown 文件形式存储。'
    'AI 根据对话内容自动加载相应 Skill，获取特定领域的操作指南。'
)

doc.add_heading('6.2 预装 Skill', level=2)
table6 = doc.add_table(rows=3, cols=3, style='Light List Accent 1')
table6.rows[0].cells[0].text = 'Skill 名称'
table6.rows[0].cells[1].text = '用途'
table6.rows[0].cells[2].text = '自动触发场景'
table6.rows[1].cells[0].text = 'system-admin'
table6.rows[1].cells[1].text = '系统包安装、卸载、查询'
table6.rows[1].cells[2].text = '用户提及安装、卸载、环境搭建等'
table6.rows[2].cells[0].text = 'web-preview'
table6.rows[2].cells[1].text = '启动开发服务器、生成预览 URL'
table6.rows[2].cells[2].text = '用户提及启动、预览、网页、服务等'

doc.add_heading('6.3 常用命令', level=2)
add_code_block(doc, '/skill list                    # 列出所有可用 Skill\n/skill load <skill-name>       # 手动加载指定 Skill')

doc.add_heading('6.4 文件位置', level=2)
table7 = doc.add_table(rows=3, cols=2, style='Light List Accent 1')
table7.rows[0].cells[0].text = '文件'
table7.rows[0].cells[1].text = '路径'
table7.rows[1].cells[0].text = 'Skill 定义文件'
table7.rows[1].cells[1].text = 'Data/opencode/.config/opencode/skills/<skill-name>/SKILL.md'
table7.rows[2].cells[0].text = '自动加载提示'
table7.rows[2].cells[1].text = 'Home/Code/opencode.json'

# ========== 7. 插件系统 ==========
doc.add_heading('7. 插件系统', level=1)
doc.add_paragraph('官方插件文档：https://opencode.ai/docs/plugins/')

doc.add_heading('7.1 安装方式一：npm 包', level=2)
doc.add_paragraph('在 ~/.config/opencode/opencode.json（全局）或项目根目录的 opencode.json 中声明：')
add_code_block(doc,
    '{\n'
    '  "$schema": "https://opencode.ai/config.json",\n'
    '  "plugin": [\n'
    '    "opencode-helicone-session",\n'
    '    "opencode-wakatime",\n'
    '    "@my-org/custom-plugin"\n'
    '  ]\n'
    '}'
)
doc.add_paragraph('启动时 OpenCode 会自动安装这些 npm 包，缓存在 ~/.cache/opencode/node_modules/。')

doc.add_heading('7.2 安装方式二：本地文件', level=2)
doc.add_paragraph('将 .js 或 .ts 文件放到插件目录下，启动时自动加载：')
doc.add_paragraph('全局插件：~/.config/opencode/plugins/', style='List Bullet')
doc.add_paragraph('项目级插件：.opencode/plugins/', style='List Bullet')

doc.add_heading('7.3 热门社区插件', level=2)
table8 = doc.add_table(rows=6, cols=2, style='Light List Accent 1')
table8.rows[0].cells[0].text = '插件名'
table8.rows[0].cells[1].text = '功能'
plugins = [
    ['opencode-wakatime', '追踪 OpenCode 使用时长'],
    ['opencode-firecrawl', '网页抓取和搜索'],
    ['oh-my-opencode', '后台 Agent、LSP/AST 工具、预置 Agent'],
    ['opencode-supermemory', '跨 Session 持久记忆'],
    ['opencode-pty', '让 AI 在 PTY 中运行后台进程并交互'],
]
for i, (name, desc) in enumerate(plugins, 1):
    table8.rows[i].cells[0].text = name
    table8.rows[i].cells[1].text = desc

# ========== 8. 使用效果展示 ==========
doc.add_heading('8. 使用效果展示', level=1)

doc.add_heading('8.1 基本对话', level=2)
doc.add_paragraph('直接用自然语言与 AI 对话，AI 会理解意图并执行操作。')
add_image_placeholder(doc, '基本对话截图')

doc.add_heading('8.2 创建文件与文件夹', level=2)
doc.add_paragraph('指令 AI 创建项目结构，如「创建一个 React 项目」。')
add_image_placeholder(doc, '创建文件夹截图')

doc.add_heading('8.3 编写代码', level=2)
doc.add_paragraph('AI 可以直接编写、修改代码文件，支持多种编程语言。')
add_image_placeholder(doc, '写代码截图')

doc.add_heading('8.4 安装环境并测试', level=2)
doc.add_paragraph('AI 可以通过 pkg-install 安装依赖，编译运行并展示结果。')
add_image_placeholder(doc, '安装环境并测试截图')

# ========== 9. TUI 模式 ==========
doc.add_heading('9. TUI 模式（可选）', level=1)
doc.add_paragraph(
    'OpenCode 官方的核心交互方式是终端 TUI，但由于 Olares 终端连接不够稳定，'
    '默认使用 Web UI。若需临时使用 TUI，可在 Control Hub 的 opencode 容器终端中执行：'
)
add_code_block(doc, 'opencode                              # 单纯启动 TUI\nopencode attach http://localhost:3000  # 启动 TUI 并连接到正在运行的 Web UI')
add_note(doc, '注：TUI 在 Olares 环境下大约能持续运行 30 分钟左右，有 Web UI 的情况下建议以 Web UI 为主。')

# ========== 10. 已知问题 ==========
doc.add_heading('10. 已知问题与解决方案', level=1)

doc.add_heading('10.1 WebUI Terminal 无法使用', level=2)
doc.add_paragraph('状态：AMD64 已修复 | ARM64 暂不支持', style='List Bullet')
p = doc.add_paragraph()
p.add_run('根因：').bold = True
p.add_run(
    'OpenCode 官方 Docker 镜像基于 Alpine Linux（musl libc），'
    '但终端功能依赖的 bun-pty 库需要 glibc。在纯 Alpine 环境中，'
    'bun-pty 无法加载原生 PTY 库，导致 Terminal 报错 '
    '"TypeError: undefined is not an object (evaluating \'lib.symbols\')"。'
)

p2 = doc.add_paragraph()
p2.add_run('解决方案：').bold = True
p2.add_run(
    '在 AMD64 下，应用安装时自动部署 glibc 兼容层并替换 OpenCode 二进制为 glibc 构建版，'
    'Terminal 功能已正常可用。ARM64 目前无 glibc 兼容方案，使用原始 musl 二进制，Terminal 不可用。'
)
doc.add_paragraph(
    '上游 Issue：#5685, #7635, #10828, #18355'
)

doc.add_heading('10.2 前端页面崩溃', level=2)
p3 = doc.add_paragraph()
p3.add_run('表现：').bold = True
p3.add_run('页面突然白屏，控制台报 TypeError: ot(...) is not a function。')

p4 = doc.add_paragraph()
p4.add_run('根因：').bold = True
p4.add_run(
    'OpenCode Web UI 的前端资源从 CDN（app.opencode.ai）实时拉取，不是打包在本地的。'
    '当 CDN 侧部署了包含 bug 的新版本前端时，所有用户会同时受影响。'
    '这属于 SolidJS 响应式框架的初始化竞态问题，是 2026 年 3 月集中爆发的一组上游 bug。'
)

p5 = doc.add_paragraph()
p5.add_run('应对措施：').bold = True
doc.add_paragraph('刷新浏览器缓存（Ctrl+Shift+R）', style='List Bullet')
doc.add_paragraph('等待 CDN 侧热修复部署', style='List Bullet')
doc.add_paragraph('关注上游 PR #16478 和 #16475 的修复进展', style='List Bullet')

doc.add_heading('10.3 WebFetch 偶发失败', level=2)
doc.add_paragraph(
    'WebFetch 功能整体可用，偶发失败通常与目标网站的可达性和当时的网络状况有关，'
    '属于正常现象，非应用自身问题。'
)

doc.add_heading('10.4 模型编辑限制', level=2)
doc.add_paragraph(
    '界面左侧的 Models 面板仅支持开关已连接的模型，不支持细节编辑。'
    '如需修改自定义 Provider 配置，需先 Disconnect 再重新连接，'
    '或直接编辑 opencode.jsonc 配置文件后重启应用。'
)

# ========== 附录 ==========
doc.add_heading('附录：Alpine Linux 包管理说明', level=1)
doc.add_paragraph(
    'OpenCode 官方镜像基于 Alpine Linux，包管理使用 apk 体系。'
    '由于 Alpine 与 Debian/Ubuntu 生态不同，部分常用工具的包名有差异。'
    '使用 pkg-install --search <关键词> 可快速查找对应的 Alpine 包名。'
    '所有 Alpine 官方仓库中的包均可通过 pkg-install 安装。'
)

# ========== 保存 ==========
output_path = '/Users/wangrongxiang/beclab/app-zips/opencode-docs/OpenCode技术使用文档.docx'
doc.save(output_path)
print(f'Document saved to: {output_path}')
