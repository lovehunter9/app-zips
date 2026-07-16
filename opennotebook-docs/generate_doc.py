"""
Generate OpenNotebook 简易使用文档.docx from the companion Markdown file.

Reads:  OpenNotebook简易使用文档.md (same directory)
Writes: OpenNotebook简易使用文档.docx (same directory)

Dependencies: python-docx
    pip install python-docx
"""

import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


HERE = os.path.dirname(os.path.abspath(__file__))
MD_PATH = os.path.join(HERE, 'OpenNotebook简易使用文档.md')
DOCX_PATH = os.path.join(HERE, 'OpenNotebook简易使用文档.docx')


# ──────────────────────────────────────────────────────────────────────────
# Document setup
# ──────────────────────────────────────────────────────────────────────────

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

for i in range(1, 5):
    hs = doc.styles[f'Heading {i}']
    hs.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    hs.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


# ──────────────────────────────────────────────────────────────────────────
# Inline markdown renderers
# ──────────────────────────────────────────────────────────────────────────

INLINE_CODE_RE = re.compile(r'`([^`]+)`')
BOLD_RE = re.compile(r'\*\*([^*]+)\*\*')
LINK_RE = re.compile(r'\[([^\]]+)\]\(([^)]+)\)')


def _split_inline(text: str):
    """Yield (kind, text) tokens: ('text'|'code'|'bold'|'link', ...)."""
    # Tokenize by finding all special spans in one pass.
    tokens = []
    i = 0
    patterns = [
        ('code', INLINE_CODE_RE),
        ('bold', BOLD_RE),
        ('link', LINK_RE),
    ]
    while i < len(text):
        best = None
        for kind, rx in patterns:
            m = rx.search(text, i)
            if m and (best is None or m.start() < best[1].start()):
                best = (kind, m)
        if best is None:
            tokens.append(('text', text[i:]))
            break
        kind, m = best
        if m.start() > i:
            tokens.append(('text', text[i:m.start()]))
        if kind == 'code':
            tokens.append(('code', m.group(1)))
        elif kind == 'bold':
            tokens.append(('bold', m.group(1)))
        elif kind == 'link':
            tokens.append(('link', m.group(1)))  # show label, drop URL
        i = m.end()
    return tokens


def _add_inline(paragraph, text: str):
    for kind, val in _split_inline(text):
        run = paragraph.add_run(val)
        if kind == 'code':
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), 'F0F0F0')
            run.element.rPr.append(shd)
        elif kind == 'bold':
            run.bold = True


# ──────────────────────────────────────────────────────────────────────────
# Block helpers
# ──────────────────────────────────────────────────────────────────────────

def add_code_block(text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x2D, 0x2D, 0x2D)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F5F5F5')
    run.element.rPr.append(shd)


def add_callout(text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.right_indent = Cm(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    _add_inline(p, text)
    for run in p.runs:
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0x8B, 0x45, 0x13)
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'FFF8E1')
        run.element.rPr.append(shd)


def add_table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers),
                      style='Light List Accent 1')
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        _add_inline(p, h)
        for run in p.runs:
            run.font.bold = True
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            cell = t.rows[r_idx].cells[c_idx]
            cell.text = ''
            _add_inline(cell.paragraphs[0], val)
    doc.add_paragraph()  # spacer


def add_bullet(text: str, level: int = 0):
    style_name = 'List Bullet' if level == 0 else f'List Bullet {min(level+1, 3)}'
    try:
        p = doc.add_paragraph(style=style_name)
    except KeyError:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Cm(0.75 * (level + 1))
    _add_inline(p, text)


def add_numbered(text: str, level: int = 0):
    try:
        p = doc.add_paragraph(style='List Number')
    except KeyError:
        p = doc.add_paragraph()
    if level > 0:
        p.paragraph_format.left_indent = Cm(0.75 * (level + 1))
    _add_inline(p, text)


def add_paragraph(text: str):
    p = doc.add_paragraph()
    _add_inline(p, text)


# ──────────────────────────────────────────────────────────────────────────
# Very small Markdown parser (tailored for this document)
# ──────────────────────────────────────────────────────────────────────────

def render_markdown(md: str):
    lines = md.split('\n')
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # HTML comment (skip)
        if stripped.startswith('<!--'):
            while i < n and '-->' not in lines[i]:
                i += 1
            i += 1
            continue

        # Horizontal rule
        if stripped == '---':
            p = doc.add_paragraph()
            run = p.add_run('─' * 40)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            i += 1
            continue

        # Headings
        m = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            title = m.group(2).strip()
            if level == 1:
                doc.add_heading(title, level=0)
            else:
                doc.add_heading(title, level=level - 1)
            i += 1
            continue

        # Fenced code block
        if stripped.startswith('```'):
            i += 1
            buf = []
            while i < n and not lines[i].lstrip().startswith('```'):
                buf.append(lines[i])
                i += 1
            add_code_block('\n'.join(buf))
            i += 1  # skip closing ```
            continue

        # Blockquote / callout
        if stripped.startswith('> '):
            buf = []
            while i < n and lines[i].lstrip().startswith('> '):
                buf.append(lines[i].lstrip()[2:])
                i += 1
            add_callout(' '.join(buf))
            continue

        # Table
        if stripped.startswith('|') and i + 1 < n and re.match(r'^\|[\s:\-|]+\|$', lines[i + 1].strip()):
            header_line = stripped
            headers = [c.strip() for c in header_line.strip('|').split('|')]
            i += 2  # skip header + separator
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                rows.append(cells)
                i += 1
            add_table(headers, rows)
            continue

        # Bullets (- or *)
        m = re.match(r'^(\s*)[-*]\s+(.*)$', line)
        if m:
            indent = len(m.group(1)) // 2
            add_bullet(m.group(2), level=indent)
            i += 1
            continue

        # Numbered list
        m = re.match(r'^(\s*)\d+\.\s+(.*)$', line)
        if m:
            indent = len(m.group(1)) // 2
            add_numbered(m.group(2), level=indent)
            i += 1
            continue

        # Plain paragraph
        add_paragraph(stripped)
        i += 1


# ──────────────────────────────────────────────────────────────────────────
# Run
# ──────────────────────────────────────────────────────────────────────────

def main():
    with open(MD_PATH, 'r', encoding='utf-8') as f:
        md = f.read()
    render_markdown(md)
    doc.save(DOCX_PATH)
    print(f'Wrote: {DOCX_PATH}')


if __name__ == '__main__':
    main()
